"""Secure PDF/PPTX storage and text extraction."""

from hashlib import sha256
from pathlib import Path
from uuid import uuid4

from fastapi import UploadFile

from app.config import get_settings


class MaterialUploadError(ValueError):
    """Raised when an upload cannot be accepted or parsed safely."""


SUPPORTED_MATERIAL_EXTENSIONS = {".pdf", ".pptx"}


def validate_material_filename(filename: str | None) -> str:
    """Return the normalized extension or reject an unsupported lesson file."""
    safe_filename = Path(filename or "lesson").name
    extension = Path(safe_filename).suffix.lower()
    if extension not in SUPPORTED_MATERIAL_EXTENSIONS:
        raise MaterialUploadError("Only PDF and PPTX lesson files are supported.")
    return extension


def _storage_path(teacher_id: str, material_id: str, extension: str) -> Path:
    root = Path(get_settings().material_storage_dir).resolve()
    path = (root / teacher_id / f"{material_id}{extension}").resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def parse_material(path: Path, media_type: str, source_id: str) -> list[dict]:
    """Extract textual page or slide blocks without claiming image understanding."""
    blocks: list[dict] = []
    if media_type == "pdf":
        from pypdf import PdfReader
        for page_number, page in enumerate(PdfReader(str(path)).pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                blocks.append({"type": "pdf_page", "page": page_number, "sourceId": source_id, "text": text})
    elif media_type == "pptx":
        from pptx import Presentation
        for slide_number, slide in enumerate(Presentation(str(path)).slides, start=1):
            texts = [shape.text.strip() for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
            if texts:
                blocks.append({"type": "pptx_slide", "slide": slide_number, "page": slide_number, "sourceId": source_id, "title": texts[0], "text": "\n".join(texts)})
    if not blocks:
        raise MaterialUploadError("No extractable text was found in the uploaded material.")
    return blocks


async def save_and_parse_upload(file: UploadFile, teacher_id: str) -> tuple[str, str, str, str, list[dict]]:
    """Validate, store, and parse a small teacher-owned lesson file."""
    filename = Path(file.filename or "lesson").name
    extension = validate_material_filename(filename)
    content = await file.read()
    if not content or len(content) > get_settings().upload_max_mb * 1024 * 1024:
        raise MaterialUploadError("The uploaded file is empty or exceeds the configured size limit.")
    media_type = extension[1:]
    material_id = str(uuid4())
    path = _storage_path(teacher_id, material_id, extension)
    path.write_bytes(content)
    try:
        blocks = parse_material(path, media_type, material_id)
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return material_id, filename, media_type, sha256(content).hexdigest(), blocks
